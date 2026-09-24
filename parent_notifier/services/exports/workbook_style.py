"""Shared look of downloaded sheets, and the guard against formula injection."""

import io
import re
from typing import TYPE_CHECKING

# openpyxl is imported where a sheet is built, so starting the app stays quick.
if TYPE_CHECKING:
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

IDENTITY_HEADERS = ["Enrollment No", "Student Name", "Parent Name", "Parent Phone"]
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")
# Phone numbers such as "+91 90000 00101" cannot hold a formula that does anything.
_PHONE_LIKE = re.compile(r"[+\-]?[\d\s()\-]+")
_HEADER_COLOR = "E4EFF2"
_WIDTHS = {1: 18, 2: 26, 3: 26, 4: 16}
SUBJECT_WIDTH = 34


def safe(value: object) -> object:
    """Text that Excel would run as a formula gets a leading apostrophe (OWASP CSV and
    formula injection guidance), so opening a downloaded sheet never runs anything."""
    if (
        isinstance(value, str)
        and value.startswith(_FORMULA_STARTS)
        and not _PHONE_LIKE.fullmatch(value)
    ):
        return f"'{value}"
    return value


def new_sheet(title: str, headers: list[str]) -> tuple["Workbook", "Worksheet"]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append([safe(header) for header in headers])
    for column, _ in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", start_color=_HEADER_COLOR)
        cell.alignment = Alignment(vertical="center")
        width = _WIDTHS.get(column, SUBJECT_WIDTH)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "B2"
    return workbook, sheet


def add_row(sheet: "Worksheet", values: list[object]) -> None:
    sheet.append([safe(value) for value in values])


def to_bytes(workbook: "Workbook") -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

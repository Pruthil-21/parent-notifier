import io

import pytest
from openpyxl import load_workbook

from parent_notifier.services.exports.prefilled_sheet import prefilled_sheet
from parent_notifier.services.exports.sample_format import sample_format
from parent_notifier.services.exports.workbook_style import safe
from parent_notifier.services.imports.apply import apply_import
from parent_notifier.services.imports.sheet_parser import parse_sheet
from parent_notifier.services.imports.sheet_reader import read_sheet
from tests.factories.academics import make_class, make_semester, make_student
from tests.factories.accounts import make_mentor
from tests.factories.workbooks import sheet_rows


def _rows(data: bytes) -> list[list[object]]:
    workbook = load_workbook(io.BytesIO(data))
    return [list(row) for row in workbook.active.iter_rows(values_only=True)]


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ('=HYPERLINK("http://x","click")', '\'=HYPERLINK("http://x","click")'),
        ("@SUM(A1)", "'@SUM(A1)"),
        ("-2+3", "'-2+3"),
        ("\t=1", "'\t=1"),
        ("+91 90000 00101", "+91 90000 00101"),
        ("Avi Shah", "Avi Shah"),
        (86, 86),
    ],
)
def test_formulas_are_neutralised_but_phones_are_not(value, expected):
    assert safe(value) == expected


def test_sheet_format_parses_cleanly():
    grid = read_sheet("format.xlsx", sample_format())
    sheet = parse_sheet(grid, 20)
    assert sheet.errors == []
    assert sheet.subjects == ["Data Structures", "DBMS", "Mathematics"]
    assert len(sheet.rows) == 1


@pytest.mark.usefixtures("app_context")
def test_prefilled_sheet_round_trips_through_import():
    mentor = make_mentor()
    class_group = make_class(mentor)
    semester = make_semester(class_group, 4)
    original = parse_sheet([[str(v) for v in row] for row in sheet_rows()], 20)
    apply_import(class_group, semester, mentor.id, "s.xlsx", original, update_identity=False)
    exported = parse_sheet(read_sheet("again.xlsx", prefilled_sheet(class_group, semester)), 20)
    assert exported.errors == []
    assert exported.subjects == original.subjects
    assert [row.cells for row in exported.rows] == [row.cells for row in original.rows]


@pytest.mark.usefixtures("app_context")
def test_empty_semester_lists_active_students_and_escapes_names():
    class_group = make_class(make_mentor())
    semester = make_semester(class_group, 1)
    make_student(class_group, "23CE002", full_name="=cmd|' /C calc'!A0")
    make_student(class_group, "23CE001", full_name="Avi Shah")
    make_student(class_group, "23CE003", full_name="Gone", status="left")
    rows = _rows(prefilled_sheet(class_group, semester))
    assert rows[0] == ["Enrollment No", "Student Name", "Parent Name", "Parent Phone"]
    assert [row[0] for row in rows[1:]] == ["23CE001", "23CE002"]
    assert rows[2][1] == "'=cmd|' /C calc'!A0"
